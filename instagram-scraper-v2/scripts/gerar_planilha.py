"""Monta uma planilha Excel estilizada (cores do Instagram) a partir dos 4 JSONs.

Lê de <perfil>/Logs_json/ e salva <perfil>/Followers_Docs/<perfil>.xlsx, com uma
aba por JSON: Perfil, Seguidores, Seguindo e De-Para.
"""

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, GradientFill, PatternFill, Side
from openpyxl.utils import get_column_letter

# Paleta do Instagram.
ROXO = "833AB4"
ROSA = "E1306C"
LARANJA = "F77737"
AMARELO = "FCAF45"
AZUL = "405DE6"
GRADIENTE = ["FCAF45", "F77737", "E1306C", "C13584", "833AB4", "405DE6"]

BRANCO = "FFFFFF"
ROSA_CLARO = "FDF0F5"   # linhas pares
TEXTO = "262626"        # cinza escuro do Instagram
CINZA = "8E8E8E"

FONTE = "Segoe UI"
BORDA_COR = "F0D5E0"

fina = Side(style="thin", color=BORDA_COR)
BORDA = Border(left=fina, right=fina, top=fina, bottom=fina)


def _fonte(tam=11, cor=TEXTO, negrito=False):
    return Font(name=FONTE, size=tam, color=cor, bold=negrito)


def _preencher(cor):
    return PatternFill("solid", fgColor=cor)


def _banner(ws, titulo, subtitulo, ncols):
    """Faixa do título com o degradê do Instagram + linha de subtítulo."""
    ultima = get_column_letter(ncols)
    ws.merge_cells(f"A1:{ultima}1")
    cel = ws["A1"]
    cel.value = titulo
    cel.font = _fonte(20, BRANCO, negrito=True)
    cel.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    cel.fill = GradientFill(degree=0, stop=GRADIENTE)
    ws.row_dimensions[1].height = 46

    ws.merge_cells(f"A2:{ultima}2")
    sub = ws["A2"]
    sub.value = subtitulo
    sub.font = _fonte(10, CINZA)
    sub.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 20


def _cabecalho(ws, colunas, linha=3):
    for i, titulo in enumerate(colunas, 1):
        cel = ws.cell(row=linha, column=i, value=titulo)
        cel.font = _fonte(11, BRANCO, negrito=True)
        cel.fill = _preencher(ROXO)
        cel.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        cel.border = BORDA
    ws.row_dimensions[linha].height = 26


def _linhas(ws, dados, larguras, inicio=4, centralizar=()):
    for n, linha in enumerate(dados):
        r = inicio + n
        fundo = _preencher(ROSA_CLARO) if n % 2 else _preencher(BRANCO)
        for c, valor in enumerate(linha, 1):
            cel = ws.cell(row=r, column=c, value=valor)
            cel.font = _fonte(10)
            cel.fill = fundo
            cel.border = BORDA
            meio = "center" if c in centralizar else "left"
            cel.alignment = Alignment(horizontal=meio, vertical="center", indent=0 if meio == "center" else 1)
    for i, larg in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = larg
    fim = inicio + len(dados) - 1
    if fim >= inicio:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(larguras))}{fim}"
    ws.freeze_panes = f"A{inicio}"
    ws.sheet_view.showGridLines = False


def _sim_nao(v):
    return "Sim" if v else ("Nao" if v is not None else "")


def _aba_pessoas(wb, nome, titulo, pessoas, cor_aba):
    ws = wb.create_sheet(nome)
    ws.sheet_properties.tabColor = cor_aba
    _banner(ws, titulo, f"{len(pessoas)} contas · gerado em {datetime.now():%d/%m/%Y %H:%M}", 5)
    _cabecalho(ws, ["#", "Usuario", "Nome", "Privado", "Verificado"])
    linhas = [[i, "@" + p.get("usuario", ""), p.get("nome") or "", _sim_nao(p.get("privado")),
               _sim_nao(p.get("verificado"))] for i, p in enumerate(pessoas, 1)]
    _linhas(ws, linhas, [6, 26, 34, 11, 13], centralizar=(1, 4, 5))
    return ws


def _aba_perfil(wb, hist):
    ws = wb.create_sheet("Perfil")
    ws.sheet_properties.tabColor = ROSA
    _banner(ws, f"@{hist.get('usuario', '')}", hist.get("nome") or "", 2)

    campos = [
        ("Nome", hist.get("nome")),
        ("Biografia", (hist.get("biografia") or "").replace("\n", "  ")),
        ("Seguidores", hist.get("seguidores")),
        ("Seguindo", hist.get("seguindo")),
        ("Posts", hist.get("posts_count")),
        ("Conta privada", _sim_nao(hist.get("privado"))),
        ("Verificada", _sim_nao(hist.get("verificado"))),
        ("Link externo", hist.get("link_externo") or ""),
        ("ID", hist.get("id")),
        ("Ultima coleta", hist.get("data_coleta")),
    ]
    r = 4
    for rotulo, valor in campos:
        a = ws.cell(row=r, column=1, value=rotulo)
        a.font = _fonte(10, BRANCO, negrito=True)
        a.fill = _preencher(ROXO)
        a.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        a.border = BORDA
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)  # valor ocupa B:D
        b = ws.cell(row=r, column=2, value=valor)
        b.font = _fonte(10)
        for c in range(2, 5):
            cel = ws.cell(row=r, column=c)
            cel.fill = _preencher(ROSA_CLARO if (r % 2) else BRANCO)
            cel.border = BORDA
        b.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        r += 1
    ws.column_dimensions["A"].width = 18
    for col, larg in (("B", 16), ("C", 14), ("D", 12)):
        ws.column_dimensions[col].width = larg

    historico = hist.get("historico") or []
    if historico:
        r += 1
        tit = ws.cell(row=r, column=1, value="Historico de coletas")
        tit.font = _fonte(13, ROXO, negrito=True)
        r += 1
        for i, col in enumerate(["Data", "Seguidores", "Seguindo", "Posts"], 1):
            cel = ws.cell(row=r, column=i, value=col)
            cel.font = _fonte(11, BRANCO, negrito=True)
            cel.fill = _preencher(ROSA)
            cel.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            cel.border = BORDA
        for n, h in enumerate(historico):
            r += 1
            fundo = _preencher(ROSA_CLARO if n % 2 else BRANCO)
            for c, chave in enumerate(["data_coleta", "seguidores", "seguindo", "posts_count"], 1):
                cel = ws.cell(row=r, column=c, value=h.get(chave))
                cel.font = _fonte(10)
                cel.fill = fundo
                cel.border = BORDA
                cel.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False


def _aba_depara(wb, dp):
    ws = wb.create_sheet("De-Para")
    ws.sheet_properties.tabColor = ROXO
    _banner(ws, "De-Para de seguidores", f"@{dp.get('perfil', '')} · {dp.get('data_coleta', '')}", 5)

    resumo = [
        ("Total de seguidores (listados)", dp.get("total_seguidores")),
        ("Total seguindo (listados)", dp.get("total_seguindo")),
        ("Nao seguem de volta", dp.get("quantidade_nao_seguem_de_volta")),
        ("Voce nao segue de volta", dp.get("quantidade_voce_nao_segue_de_volta")),
        ("Listas completas", _sim_nao(dp.get("listas_completas"))),
    ]
    r = 4
    for rotulo, valor in resumo:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)  # rótulo em A:C
        a = ws.cell(row=r, column=1, value=rotulo)
        a.font = _fonte(10, BRANCO, negrito=True)
        for c in range(1, 4):
            cel = ws.cell(row=r, column=c)
            cel.fill = _preencher(ROSA)
            cel.border = BORDA
        a.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)  # valor em D:E
        b = ws.cell(row=r, column=4, value=valor)
        b.font = _fonte(11, ROXO, negrito=True)
        b.alignment = Alignment(horizontal="center", vertical="center")
        for c in (4, 5):
            ws.cell(row=r, column=c).border = BORDA
        r += 1

    for aviso in dp.get("avisos") or []:
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        cel = ws.cell(row=r, column=1, value="⚠ " + aviso)
        cel.font = _fonte(9, LARANJA)
        cel.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        ws.row_dimensions[r].height = 30

    def tabela(r, titulo, pessoas):
        """Título + tabela de pessoas a partir da linha r. Devolve a próxima linha livre."""
        tit = ws.cell(row=r, column=1, value=titulo)
        tit.font = _fonte(13, ROXO, negrito=True)
        r += 1
        _cabecalho(ws, ["#", "Usuario", "Nome", "Privado", "Verificado"], linha=r)
        for n, p in enumerate(pessoas):
            rr = r + 1 + n
            fundo = _preencher(ROSA_CLARO if n % 2 else BRANCO)
            for c, valor in enumerate([n + 1, "@" + p.get("usuario", ""), p.get("nome") or "",
                                       _sim_nao(p.get("privado")), _sim_nao(p.get("verificado"))], 1):
                cel = ws.cell(row=rr, column=c, value=valor)
                cel.font = _fonte(10)
                cel.fill = fundo
                cel.border = BORDA
                meio = "center" if c in (1, 4, 5) else "left"
                cel.alignment = Alignment(horizontal=meio, vertical="center", indent=0 if meio == "center" else 1)
        return r + 1 + len(pessoas) + 2  # pula 2 linhas até a próxima tabela

    r += 2
    r = tabela(r, "Seguidos que NAO retribuem (candidatos a unfollow)", dp.get("nao_seguem_de_volta") or [])
    tabela(r, "Te seguem e voce NAO segue de volta", dp.get("voce_nao_segue_de_volta") or [])
    for i, larg in enumerate([6, 26, 34, 11, 13], 1):
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.sheet_view.showGridLines = False


def _ler(caminho, padrao):
    try:
        return json.loads(Path(caminho).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return padrao


def gerar_planilha(pasta_logs, pasta_docs, perfil):
    """Monta a planilha a partir dos JSONs em pasta_logs. Devolve o caminho ou None."""
    pasta_logs, pasta_docs = Path(pasta_logs), Path(pasta_docs)
    hist = _ler(pasta_logs / "historico_perfil.json", {})
    seguidores = _ler(pasta_logs / "seguidores.json", [])
    seguindo = _ler(pasta_logs / "seguindo.json", [])
    depara = _ler(pasta_logs / "unfollowers_followed.json", {})
    if not any([hist, seguidores, seguindo, depara]):
        return None

    wb = Workbook()
    wb.remove(wb.active)  # tira a aba vazia padrão
    if hist:
        _aba_perfil(wb, hist)
    if seguidores:
        _aba_pessoas(wb, "Seguidores", "Seguidores", seguidores, ROSA)
    if seguindo:
        _aba_pessoas(wb, "Seguindo", "Seguindo", seguindo, LARANJA)
    if depara:
        _aba_depara(wb, depara)
    if not wb.sheetnames:
        return None

    pasta_docs.mkdir(parents=True, exist_ok=True)
    destino = pasta_docs / f"{perfil}.xlsx"
    wb.save(destino)
    return destino


if __name__ == "__main__":
    base = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(".")
    perfil = sys.argv[2] if len(sys.argv) > 2 else base.name
    saida = gerar_planilha(base / "Logs_json", base / "Followers_Docs", perfil)
    print("Planilha gerada em" if saida else "Sem dados para a planilha.", saida or "")
